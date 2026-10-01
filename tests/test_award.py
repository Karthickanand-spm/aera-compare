"""Tests for award.py: vendor sets, discounts, label toggles and confirmation. No API calls."""

import io

import openpyxl
import pandas as pd
import pytest

from aera.analyst import AnalystData
from aera.award import (
    CAPPED, CHEAPEST, SINGLE, AwardSettings, build_award, check_discount, confirm_award,
    confirm_blockers, order_value_threshold, award_to_excel, vendor_choices,
)

QTY = 1000


def _row(line, vendor, price, label="Comparable", needs_review=False, confirmed=False):
    counted = label in ("Comparable", "Comparable with assumption")
    return {"rfx_line_id": line, "vendor": vendor.upper(), "display_name": vendor,
            "price_inr_per_piece": price, "label": label, "included_in_totals": counted,
            "needs_review": needs_review, "buyer_confirmed": confirmed, "assumptions": [],
            "confidence": "low" if needs_review else "high", "confidence_reasons": ["test"],
            "raw_price_text": f"{price}", "source_file": f"{vendor}.xlsx", "source_snippet": f"{price}",
            "page": None}


def _data(rows, lines=(1, 2, 3), last_year=None) -> AnalystData:
    df = pd.DataFrame(rows)
    rfx_lines = pd.DataFrame([{"line_id": i, "description": f"Box {i}", "annual_qty": QTY, "uom": "pcs",
                               "nominal_weight_g": 200.0} for i in lines])
    ly = pd.DataFrame(sorted((last_year or {}).items()), columns=["line_id", "price_inr_per_piece"])
    return AnalystData(df, pd.DataFrame(), rfx_lines, ly, {"USD": 94.5}, "2026-09-25, illustrative rate")


def _summary(names, status=None) -> pd.DataFrame:
    status = status or {}
    return pd.DataFrame([{"vendor": n.upper(), "display_name": n, "quality_status": status.get(n, "PASS"),
                          "quality_reasons": [f"{status.get(n, 'PASS')}: defect rate 3% is above the 2% limit"],
                          "freight": "included", "open_risks": []} for n in names])


# ---------- Vendor sets ----------

def _three_vendors():
    # qty 1000 on each line. Pairs: A+B = 28, A+C = 27.5, B+C = 25.5 (x 1000).
    prices = {"A": (10, 10, 10), "B": (8, 12, 12), "C": (12, 8, 9.5)}
    return [_row(line, v, p[line - 1]) for v, p in prices.items() for line in (1, 2, 3)]


def test_capped_at_two_picks_the_cheapest_pair():
    data = _data(_three_vendors())
    award = build_award(data, _summary("ABC"), AwardSettings(CAPPED, ["A", "B", "C"], cap=2))
    assert award.vendors_used == ["B", "C"]
    assert award.total_inr == pytest.approx(25.5 * QTY)
    assert award.combinations_tried == 3
    assert list(award.lines["display_name"]) == ["B", "C", "C"]


def test_single_vendor_picks_the_cheapest_vendor_covering_most_lines():
    rows = _three_vendors() + [_row(1, "D", 1.0), _row(2, "D", None, "Not quoted"), _row(3, "D", None, "Not quoted")]
    award = build_award(_data(rows), _summary("ABCD"), AwardSettings(SINGLE, ["A", "B", "C", "D"]))
    # D is cheapest on paper but quoted one line only, so it must not win.
    assert award.vendors_used == ["C"]
    assert award.total_inr == pytest.approx(29.5 * QTY)


def test_line_no_allowed_vendor_priced_is_not_awarded_never_zero():
    rows = [_row(1, "A", 10), _row(2, "A", None, "Not quoted")]
    award = build_award(_data(rows, lines=(1, 2)), _summary("A"), AwardSettings(CHEAPEST, ["A"]))
    assert award.unawarded_lines == [2]
    line2 = award.lines[award.lines["rfx_line_id"] == 2].iloc[0]
    assert line2["annual_inr"] is None or pd.isna(line2["annual_inr"])
    assert award.total_inr == pytest.approx(10 * QTY)
    assert any(r["severity"] == "high" and "Line 2" in r["text"] for r in award.risks)


def test_only_quality_pass_vendors_are_allowed_by_default_with_reasons():
    choices = vendor_choices(_summary("ABC", {"B": "FAIL", "C": "UNCLEAR"})).set_index("display_name")
    assert choices.loc["A", "default_allowed"]
    assert not choices.loc["B", "default_allowed"] and not choices.loc["C", "default_allowed"]
    assert choices.loc["B", "reason"] == "defect rate 3% is above the 2% limit"


# ---------- 'Not comparable' rows ----------

def test_not_comparable_rows_are_excluded_by_default():
    rows = [_row(1, "A", 10), _row(1, "B", 7, label="Not comparable")]
    data = _data(rows, lines=(1,))
    off = build_award(data, _summary("AB"), AwardSettings(CHEAPEST, ["A", "B"]))
    assert off.vendors_used == ["A"]
    assert off.total_inr == pytest.approx(10 * QTY)

    on = build_award(data, _summary("AB"), AwardSettings(CHEAPEST, ["A", "B"], include_not_comparable=True))
    assert on.vendors_used == ["B"]
    assert on.total_inr == pytest.approx(7 * QTY)
    # The caller's data is untouched by the toggle.
    assert not data.df.loc[data.df["vendor"] == "B", "included_in_totals"].iloc[0]


# ---------- Discounts ----------

@pytest.mark.parametrize("condition, amount, inclusive", [
    ("where the annual order value placed with us exceeds Rs 25,00,000", 2_500_000, False),
    ("annual purchase value above ₹25 lakh", 2_500_000, False),
    ("minimum annual business of INR 1.2 crore", 12_000_000, True),
])
def test_order_value_threshold_is_read_from_the_condition(condition, amount, inclusive):
    assert order_value_threshold(condition) == {"amount_inr": amount, "inclusive": inclusive}


@pytest.mark.parametrize("condition", [
    "payment within 7 days",  # not an order value
    "order value above Rs 50,000 per order",  # not annual
    "annual order value above USD 30,000",  # not INR
])
def test_conditions_code_cannot_check_are_left_alone(condition):
    assert order_value_threshold(condition) is None
    result = check_discount({"text": "2% off", "condition": condition, "percent": 2}, 10_000_000, apply=True)
    assert result["status"] == "cannot check" and not result["applied"]


DISCOUNT = {"text": "10% off total invoice value", "percent": 10.0, "source_snippet": "10% off",
            "condition": "where the annual order value placed with us exceeds Rs 25,000"}


def test_discount_is_applied_only_when_its_condition_is_met():
    # A wins line 1 only (₹20,000 a year): under the ₹25,000 threshold.
    small = [_row(1, "A", 20), _row(1, "B", 25)]
    award = build_award(_data(small, lines=(1,)), _summary("AB"),
                        AwardSettings(CHEAPEST, ["A", "B"], apply_discounts=True), {"A": [DISCOUNT]})
    assert award.discounts[0]["status"] == "not met" and not award.discounts[0]["applied"]
    assert award.total_inr == pytest.approx(20 * QTY)

    # A wins two lines (₹40,000 a year): over the threshold, so 10% comes off.
    big = small + [_row(2, "A", 20), _row(2, "B", 25)]
    award = build_award(_data(big, lines=(1, 2)), _summary("AB"),
                        AwardSettings(CHEAPEST, ["A", "B"], apply_discounts=True), {"A": [DISCOUNT]})
    assert award.discounts[0]["status"] == "met" and award.discounts[0]["applied"]
    assert award.total_inr == pytest.approx(36 * QTY)
    assert award.lines["price_inr_per_piece"].tolist() == pytest.approx([18, 18])
    assert award.vendor_subtotals.iloc[0]["discount_inr"] == pytest.approx(4 * QTY)


def test_met_discount_is_not_applied_while_the_toggle_is_off():
    rows = [_row(1, "A", 20), _row(2, "A", 20)]
    award = build_award(_data(rows, lines=(1, 2)), _summary("A"), AwardSettings(CHEAPEST, ["A"]),
                        {"A": [DISCOUNT]})
    assert award.discounts[0]["status"] == "met" and not award.discounts[0]["applied"]
    assert award.total_inr == pytest.approx(40 * QTY)


# ---------- Confirmation ----------

def test_confirm_is_blocked_while_review_rows_are_open():
    open_review = [_row(1, "A", 10, needs_review=True), _row(2, "A", 12)]
    award = build_award(_data(open_review, lines=(1, 2)), _summary("A"), AwardSettings(CHEAPEST, ["A"]))
    assert confirm_blockers(award) == ["line 1 (A)"]
    with pytest.raises(ValueError, match="line 1"):
        confirm_award(award, "ok", "2026-10-01 10:00 UTC")

    reviewed = [_row(1, "A", 10, needs_review=True, confirmed=True), _row(2, "A", 12)]
    award = build_award(_data(reviewed, lines=(1, 2)), _summary("A"), AwardSettings(CHEAPEST, ["A"]))
    record = confirm_award(award, "  Go with A  ", "2026-10-01 10:00 UTC")
    assert record["vendors"] == ["A"] and record["note"] == "Go with A"
    assert record["total_inr"] == pytest.approx(22 * QTY)


def test_review_row_on_a_losing_vendor_does_not_block():
    rows = [_row(1, "A", 10), _row(1, "B", 15, needs_review=True)]
    award = build_award(_data(rows, lines=(1,)), _summary("AB"), AwardSettings(CHEAPEST, ["A", "B"]))
    assert confirm_blockers(award) == []


# ---------- Saving and export ----------

def test_saving_uses_only_lines_with_a_last_year_price():
    rows = [_row(1, "A", 9), _row(2, "A", 5)]
    award = build_award(_data(rows, lines=(1, 2), last_year={1: 10}), _summary("A"), AwardSettings(CHEAPEST, ["A"]))
    assert award.saving_inr == pytest.approx(1 * QTY)
    assert award.saving_pct == pytest.approx(10.0)
    assert award.lines_without_last_year == [2]


def test_excel_export_has_every_sheet_and_raw_rupee_numbers():
    award = build_award(_data(_three_vendors()), _summary("ABC"), AwardSettings(CHEAPEST, ["A", "B", "C"]))
    wb = openpyxl.load_workbook(io.BytesIO(award_to_excel(award, "RFX-1", "1 USD = ₹94.50 (illustrative)",
                                                          None, "2026-10-01")))
    assert wb.sheetnames == ["Award", "Vendor subtotals", "Assumptions", "Provenance", "Open risks",
                             "Freight sensitivity"]
    ws = wb["Award"]
    assert ws["B1"].value == "RFX-1" and "illustrative" in ws["B3"].value
    assert ws["B5"].value.startswith("Not confirmed")
    header = [c.value for c in ws[8]]
    cell = ws.cell(row=9, column=header.index("Annual ₹") + 1)
    assert isinstance(cell.value, (int, float)) and "₹" in cell.number_format
