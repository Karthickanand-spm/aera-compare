"""Turn a vendor reply file into a Payload ready for the Claude Messages API.

Excel, Word and email are converted to plain text here, in code. PDFs and
images are passed to Claude as base64 blocks because Claude reads them natively.
Every payload carries a sha256 of the file bytes, used as the extraction cache key.

Run as a script to preview what the loaders see:
    python -m aera.ingest data/sample/replies
"""

import base64
import email
import hashlib
import io
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from email import policy
from email.header import decode_header, make_header
from pathlib import Path

import docx
import openpyxl
from docx.table import Table
from PIL import Image, ImageOps

# Images with a long side above this many pixels are downscaled before sending.
MAX_IMAGE_SIDE_PX = 2000

KIND_BY_SUFFIX = {
    ".xlsx": "excel",
    ".xlsm": "excel",
    ".docx": "word",
    ".eml": "email",
    ".pdf": "pdf",
    ".jpg": "image",
    ".jpeg": "image",
    ".png": "image",
}

IMAGE_MEDIA_TYPES = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png"}


@dataclass(frozen=True)
class Payload:
    filename: str
    kind: str  # excel | word | email | pdf | image
    sha256: str  # hash of the original file bytes, for caching
    text: str | None  # set for excel / word / email, None for pdf / image
    content_blocks: list[dict] = field(default_factory=list)  # Messages API blocks


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def load_reply(path: str | Path) -> Payload:
    path = Path(path)
    kind = KIND_BY_SUFFIX.get(path.suffix.lower())
    if kind is None:
        supported = ", ".join(sorted(KIND_BY_SUFFIX))
        raise ValueError(f"Unsupported file type '{path.suffix}' for {path.name}. Supported: {supported}")

    data = path.read_bytes()
    digest = sha256_bytes(data)

    if kind in ("excel", "word", "email"):
        if kind == "excel":
            text = excel_to_text(data)
        elif kind == "word":
            text = word_to_text(data)
        else:
            text = email_to_text(data)
        blocks = [{"type": "text", "text": f"File: {path.name}\n\n{text}"}]
        return Payload(path.name, kind, digest, text, blocks)

    if kind == "pdf":
        block = {
            "type": "document",
            "source": {"type": "base64", "media_type": "application/pdf", "data": _b64(data)},
        }
        return Payload(path.name, kind, digest, None, [block])

    image_bytes, media_type = prepare_image(data, IMAGE_MEDIA_TYPES[path.suffix.lower()])
    block = {
        "type": "image",
        "source": {"type": "base64", "media_type": media_type, "data": _b64(image_bytes)},
    }
    return Payload(path.name, kind, digest, None, [block])


def _b64(data: bytes) -> str:
    return base64.standard_b64encode(data).decode("ascii")


# ---------- Excel ----------

def _cell_text(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def excel_to_text(data: bytes) -> str:
    """Every sheet, every non-empty row, one line per row with cell references.

    Formula cells show their last calculated value. If the file has no saved
    value for a formula, the formula itself is shown instead so nothing is lost.
    """
    values_wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    formulas_wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)

    out: list[str] = []
    for ws in values_wb.worksheets:
        fws = formulas_wb[ws.title]
        out.append(f"=== Sheet: {ws.title} ===")

        merged = sorted(ws.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col))
        if merged:
            out.append("Merged cells (the value sits in the top-left cell and spans the whole range):")
            for rng in merged:
                top_left = ws.cell(row=rng.min_row, column=rng.min_col).value
                shown = f'"{_cell_text(top_left)}"' if top_left is not None else "(empty)"
                out.append(f"  {rng.coord}: {shown}")
            out.append("")

        for row in ws.iter_rows():
            parts = []
            for cell in row:
                value = cell.value
                if value is None:
                    formula = fws[cell.coordinate].value
                    if isinstance(formula, str) and formula.startswith("="):
                        value = f"{formula} (formula, no saved value)"
                if value is None or (isinstance(value, str) and not value.strip()):
                    continue
                parts.append(f"{ws.title}!{cell.coordinate}: {_cell_text(value)}")
            if parts:
                out.append(" | ".join(parts))
        out.append("")
    return "\n".join(out).rstrip() + "\n"


# ---------- Word ----------

def word_to_text(data: bytes) -> str:
    """All paragraphs and tables in document order. Tables become pipe-separated rows."""
    document = docx.Document(io.BytesIO(data))
    out: list[str] = []
    table_no = 0
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            table_no += 1
            out.append(f"[Table {table_no}]")
            for row in item.rows:
                cells = [_dedupe_merged(row.cells, i) for i in range(len(row.cells))]
                out.append(" | ".join(c for c in cells if c is not None))
            out.append(f"[End of table {table_no}]")
        else:
            text = item.text.strip()
            if text:
                out.append(text)
    return "\n".join(out) + "\n"


def _dedupe_merged(cells, i: int) -> str | None:
    """python-docx repeats a horizontally merged cell once per grid column; show it once."""
    if i > 0 and cells[i]._tc is cells[i - 1]._tc:
        return None
    return cells[i].text.strip()


# ---------- Email ----------

def email_to_text(data: bytes) -> str:
    """From / Date / Subject headers plus the plain-text body."""
    msg = email.message_from_bytes(data, policy=policy.default)
    # Headers are read as written. The default policy would rewrite them (e.g. "fix" the
    # weekday in a Date), and the vendor's exact text matters for source snippets.
    raw = email.message_from_bytes(data)
    lines = [
        f"{h}: {make_header(decode_header(raw[h]))}"
        for h in ("From", "Date", "Subject")
        if raw[h]
    ]

    body_part = msg.get_body(preferencelist=("plain",))
    if body_part is not None:
        body = body_part.get_content()
    else:
        html_part = msg.get_body(preferencelist=("html",))
        body = _strip_html(html_part.get_content()) if html_part is not None else ""

    return "\n".join(lines) + "\n\n" + body.strip() + "\n"


def _strip_html(html: str) -> str:
    html = re.sub(r"(?is)<(script|style).*?</\1>", "", html)
    html = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>", "\n", html)
    return re.sub(r"<[^>]+>", "", html)


# ---------- Images ----------

def prepare_image(data: bytes, media_type: str) -> tuple[bytes, str]:
    """Return image bytes to send. Downscales if the long side is over MAX_IMAGE_SIDE_PX."""
    img = Image.open(io.BytesIO(data))
    if max(img.size) <= MAX_IMAGE_SIDE_PX:
        return data, media_type

    img = ImageOps.exif_transpose(img)  # phone photos: apply rotation before resizing
    img.thumbnail((MAX_IMAGE_SIDE_PX, MAX_IMAGE_SIDE_PX), Image.LANCZOS)
    buf = io.BytesIO()
    if media_type == "image/png":
        img.save(buf, format="PNG", optimize=True)
    else:
        img.convert("RGB").save(buf, format="JPEG", quality=90)
    return buf.getvalue(), media_type


# ---------- Script ----------

def _main(args: list[str]) -> int:
    if not args:
        print("Usage: python -m aera.ingest <folder or files...>")
        return 1
    sys.stdout.reconfigure(errors="replace")  # Windows consoles may not print every character

    paths: list[Path] = []
    for arg in args:
        p = Path(arg)
        paths.extend(sorted(f for f in p.iterdir() if f.is_file()) if p.is_dir() else [p])

    for p in paths:
        print("=" * 70)
        print(p.name)
        try:
            payload = load_reply(p)
        except Exception as e:  # show the problem and keep going with other files
            print(f"  ERROR: {e}")
            continue
        print(f"  kind:   {payload.kind}")
        print(f"  sha256: {payload.sha256}")
        if payload.text is not None:
            print("  text (first 500 chars):")
            print(payload.text[:500])
        else:
            print(f"  binary: {payload.kind}")
    return 0


if __name__ == "__main__":
    sys.exit(_main(sys.argv[1:]))
